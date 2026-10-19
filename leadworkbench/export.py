"""Safe business exports (design §5). Refusal records are excluded.

All companies of the current filter are selected once into temp.export_ids, then
every kind of data is read with one batched query (no per-company queries).

- layout="companies" (API default): the established Companies/Contacts/Locations/
  Sources tables; XLSX Contacts/Locations/Sources append a company_name column.
- layout="entries": one row per contact entry, Chinese labels; XLSX adds 企业/地点/
  来源证据 sheets, each starting with 企业名称.
"""
import csv
import io
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .address_city import city_from_address
from .errors import Invalid
from .contact_fields import _source
from .contact_projection import categorize_outreach, group_contact_records, merge_versions
from .contact_tiers import TIER_LABEL, contact_tier
from .labels import STATUS_LABELS, business_type_label

UNUSABLE = {"pending", "reject", "invalid", "superseded"}
COMPANY_HEADERS = ["id", "name", "address", "phone", "website", "city", "country", "business_types", "status",
                   "source_kind", "notes", "created_at", "last_checked_at", "source_url", "observed_at"]
MARKET_HEADERS = ["market_segment", "market_changed", "first_seen", "last_seen"]
CONTACT_KEYS = ["company_id", "name", "role", "email", "url", "kind", "role_status",
                "association_status", "review_status", "source_url", "evidence", "observed_at"]
LOCATION_KEYS = ["company_id", "address", "latitude", "longitude", "source_url", "observed_at"]
SOURCE_KEYS = ["company_id", "field", "value", "source_url", "evidence", "observed_at", "method", "query"]
ENTRY_HEADERS = ["企业名称", "企业编号", "实际城市", "城市依据", "目标市场", "来源类别", "业务类型", "客户匹配状态",
                 "入口等级", "渠道", "联系方式", "其他联系方式", "联系人", "岗位", "来源", "来源链接", "采集时间",
                 "审核状态", "企业官网"]
ENTRY_COMPANY_HEADERS = ["企业名称", "企业编号", "地址", "电话", "企业官网", "实际城市", "城市依据", "目标市场",
                         "来源类别", "业务类型", "客户匹配状态", "最佳入口等级", "联系入口数（已去重）", "缺口原因",
                         "来源链接", "采集时间"]
ENTRY_LOCATION_HEADERS = ["企业名称", "企业编号", "地址", "纬度", "经度", "来源链接", "采集时间"]
ENTRY_SOURCE_HEADERS = ["企业名称", "企业编号", "字段", "值", "来源链接", "来源原文", "采集时间", "方式", "检索词"]
METHOD_LABELS = {"maps": "Google Maps", "website": "企业官网", "manual": "人工录入", "social": "社媒"}


FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def safe(value):
    """Plain cell text. The formula guard is applied per format in _write."""
    if value is None:
        return ""
    if isinstance(value, list):
        value = ", ".join(str(v) for v in value)
    return str(value)


def _formula_like(value):
    return isinstance(value, str) and value.lstrip().startswith(FORMULA_START)


def _csv_cell(value):
    # CSV has no cell types: a leading ' keeps spreadsheets from evaluating it.
    return "'" + value if _formula_like(value) else value


def _xlsx_cell(ws, value):
    # XLSX stores these as explicit text cells (no visible apostrophe); openpyxl
    # would otherwise turn a leading "=" into a formula.
    if not _formula_like(value):
        return value
    cell = WriteOnlyCell(ws, value=value)
    cell.data_type, cell.number_format = "s", "@"
    return cell


def _batch(con, sql, params=()):
    return con.execute(sql.replace("{ids}", "JOIN temp.export_ids e ON e.id="), params).fetchall()


def _collect(store, con, filters, *, observations):
    query = store._company_query(con, filters.get("q", ""), filters.get("status", ""), filters.get("type", ""),
                                 filters.get("market_id", ""), filters.get("run_id", ""), filters.get("segment", ""),
                                 filters.get("contact", ""), filters.get("category", ()),
                                 filters.get("excluded", "hide"), filters.get("best_tier", ""))
    keyed = store._company_ids(con, query, filters.get("sort", "newest"))
    companies = [c for c in store._company_rows(con, keyed, query) if c["status"] != "do_not_contact"]
    con.execute("CREATE TEMP TABLE IF NOT EXISTS export_ids(id TEXT PRIMARY KEY, ord INTEGER)")
    con.execute("DELETE FROM temp.export_ids")
    con.executemany("INSERT INTO temp.export_ids(id,ord) VALUES(?,?)", [(c["id"], i) for i, c in enumerate(companies)])
    data = {"companies": companies, "market_id": query["market_id"], "latest": {}, "contacts": {}, "versions": {},
            "locations": {}, "observations": {}, "listings": {}, "markets": {}, "phone_obs": {}}
    for r in _batch(con, """SELECT company_id,source_url,observed_at FROM (SELECT o.company_id,o.source_url,o.observed_at,
            ROW_NUMBER() OVER (PARTITION BY o.company_id ORDER BY o.observed_at DESC,o.rowid DESC) rn
            FROM observations o {ids}o.company_id) WHERE rn=1"""):
        data["latest"][r[0]] = {"source_url": r[1], "observed_at": r[2]}
    for r in _batch(con, """SELECT ct.*,mc.phone manual_phone,mc.contact_id IS NOT NULL manual,d.adoption_status
            FROM contacts ct {ids}ct.company_id LEFT JOIN manual_contacts mc ON mc.contact_id=ct.id
            LEFT JOIN contact_decisions d ON d.contact_id=ct.id ORDER BY e.ord,ct.rowid"""):
        data["contacts"].setdefault(r["company_id"], []).append(store._contact_row(r))
    for r in _batch(con, """SELECT v.contact_id,v.payload FROM contact_versions v JOIN contacts ct ON ct.id=v.contact_id
            {ids}ct.company_id ORDER BY v.contact_id,v.observed_at,v.rowid"""):
        data["versions"].setdefault(r[0], []).append(json.loads(r[1]))
    for r in _batch(con, "SELECT l.* FROM locations l {ids}l.company_id ORDER BY e.ord,l.rowid"):
        data["locations"].setdefault(r["company_id"], []).append(dict(r))
    if observations:
        for r in _batch(con, "SELECT o.* FROM observations o {ids}o.company_id ORDER BY e.ord,o.observed_at,o.rowid"):
            data["observations"].setdefault(r["company_id"], []).append(dict(r))
    for r in _batch(con, "SELECT l.company_id,l.source_url FROM listings l {ids}l.company_id ORDER BY l.last_seen DESC"):
        data["listings"].setdefault(r[0], r[1])
    for r in _batch(con, """SELECT company_id,source_url,method,observed_at FROM (SELECT o.company_id,o.source_url,o.method,
            o.observed_at,ROW_NUMBER() OVER (PARTITION BY o.company_id ORDER BY o.observed_at DESC,o.rowid DESC) rn
            FROM observations o {ids}o.company_id WHERE o.field='phone') WHERE rn=1"""):
        data["phone_obs"][r[0]] = dict(r)
    for r in _batch(con, """SELECT m.company_id,t.name FROM market_memberships m {ids}m.company_id
            JOIN tasks t ON t.id=m.market_id ORDER BY m.first_seen,m.market_id"""):
        data["markets"].setdefault(r[0], []).append(r[1])
    if data["market_id"]:
        row = con.execute("SELECT name FROM tasks WHERE id=?", (data["market_id"],)).fetchone()
        data["market_name"] = row[0] if row else ""
    return data


def _grouped_contacts(data, company):
    """Usable contacts of one company, grouped exactly like the detail page."""
    records = [row for row in merge_versions(data["contacts"].get(company["id"], []), data["versions"])
               if row.get("review_status") not in UNUSABLE and row.get("adoption_status") != "excluded"]
    return [row for row in categorize_outreach(group_contact_records(records, company.get("website", "")))
            if row.get("outreach_category") != "irrelevant"]


class ExportCancelled(Exception):
    """should_stop() became true; the caller removes the partial file."""


PROGRESS_COMPANIES = 200
PROGRESS_ROWS = 1000


class _Progress:
    """Progress/cancel hooks: every 200 companies while building, every 1000 rows while writing."""
    def __init__(self, progress=None, should_stop=None):
        self.progress, self.should_stop = progress, should_stop
        self.rows = 0

    def report(self, phase, **values):
        if self.should_stop and self.should_stop():
            raise ExportCancelled()
        if self.progress:
            self.progress(phase, **values)


def normalize_filters(store, layout, filters):
    """run_id -> market_id and layout check; raises NotFound/Invalid before any work starts."""
    filters = dict(filters)
    if filters.get("run_id") and not filters.get("market_id"):
        market_id = store.run(filters["run_id"]).get("task_id")
        if market_id:
            filters["market_id"] = market_id
    if layout not in {"companies", "entries"}:
        raise Invalid("无效导出布局")
    return filters


def export_to_file(store, target, format="csv", layout="companies", filters=None, *, progress=None,
                   should_stop=None):
    """Write one export to ``target`` (a path or a binary file object).

    The data is read in one deferred (read-only) transaction, so the file is a
    consistent snapshot and no write lock is held while reading."""
    filters = normalize_filters(store, layout, filters or {})
    hooks = _Progress(progress, should_stop)
    store._refresh_facets()
    hooks.report("reading")
    with store.connect() as con:
        con.execute("BEGIN")
        try:
            data = _collect(store, con, filters, observations=format == "xlsx")
        finally:
            con.rollback()
    hooks.report("building", processed=0, total=len(data["companies"]))
    if layout == "entries":
        sheets = _entries_layout(data, format, hooks)
    else:
        sheets = _companies_layout(data, format, bool(filters.get("market_id")), hooks)
    _write(format, sheets, target, hooks)
    return {"total": len(data["companies"]), "rows": hooks.rows}


def export_data(store, format="csv", layout="companies", **filters):
    buffer = io.BytesIO()
    export_to_file(store, buffer, format, layout, filters)
    return buffer.getvalue()


def _write(format, sheets, target, hooks=None):
    """sheets: [(title, headers, rows)]. CSV writes only the first sheet."""
    hooks = hooks or _Progress()
    if format == "csv":
        stream = io.StringIO(newline="")
        writer = csv.writer(stream)
        title, headers, rows = sheets[0]
        writer.writerow(headers)
        for start in range(0, len(rows), PROGRESS_ROWS):
            writer.writerows([_csv_cell(v) for v in row] for row in rows[start:start + PROGRESS_ROWS])
            hooks.rows = min(len(rows), start + PROGRESS_ROWS)
            hooks.report("writing", rows=hooks.rows)
        data = ("\ufeff" + stream.getvalue()).encode("utf-8")
        if isinstance(target, (str, Path)):
            Path(target).write_bytes(data)
        else:
            target.write(data)
        return
    book = Workbook(write_only=True)
    font, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="173C46")
    for title, headers, rows in sheets:
        ws = book.create_sheet(title)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        for index, header in enumerate(headers, 1):
            width = max([len(str(header))] + [len(str(row[index - 1] or "")) for row in rows[:40]]) + 2
            ws.column_dimensions[get_column_letter(index)].width = min(55, max(14, width))
        cells = []
        for header in headers:
            cell = WriteOnlyCell(ws, value=header)
            cell.font, cell.fill = font, fill
            cells.append(cell)
        ws.append(cells)
        for row in rows:
            ws.append([_xlsx_cell(ws, v) for v in row])
            hooks.rows += 1
            if hooks.rows % PROGRESS_ROWS == 0:
                hooks.report("writing", rows=hooks.rows)
    hooks.report("writing", rows=hooks.rows)
    book.save(target)
def _companies_layout(data, format, by_market, hooks=None):
    hooks = hooks or _Progress()
    headers = COMPANY_HEADERS + (MARKET_HEADERS if by_market else [])
    company_rows = []
    contacts, locations, sources = [], [], []
    total = len(data["companies"])
    for index, company in enumerate(data["companies"]):
        if index and index % PROGRESS_COMPANIES == 0:
            hooks.report("building", processed=index, total=total)
        values = {**company, **data["latest"].get(company["id"], {"source_url": "", "observed_at": ""})}
        company_rows.append([safe(values.get(k)) for k in headers])
        if format == "csv":
            continue
        name = safe(company["name"])
        for row in _grouped_contacts(data, company):
            contacts.append([safe(row.get(k)) for k in CONTACT_KEYS] + [name])
            for source in row["sources"]:
                sources.append([safe({'company_id': company['id'], 'field': 'contact',
                                      'value': row.get('email') or row.get('name') or row.get('url'),
                                      'method': 'website', **source}.get(k)) for k in SOURCE_KEYS] + [name])
        for row in data["locations"].get(company["id"], []):
            locations.append([safe(row[k]) for k in LOCATION_KEYS] + [name])
        for row in data["observations"].get(company["id"], []):
            sources.append([safe(row[k]) for k in SOURCE_KEYS] + [name])
    hooks.report("building", processed=total, total=total)
    return [("Companies", headers, company_rows),
            ("Contacts", CONTACT_KEYS + ["company_name"], contacts),
            ("Locations", LOCATION_KEYS + ["company_name"], locations),
            ("Sources", SOURCE_KEYS + ["company_name"], sources)]


def _company_columns(data, company):
    city = city_from_address(company.get("address"), company.get("country"))
    market = data.get("market_name") if data["market_id"] else "、".join(data["markets"].get(company["id"], []))
    types = "、".join(business_type_label(t) for t in company.get("business_types") or []) or "、".join(
        dict.fromkeys(business_type_label(h["business_type"], hint=True) for h in company.get("type_hints") or []))
    return dict(city=city or company.get("city") or "", basis="地址" if city else "目标市场（地址未解析）",
                market=market or "", categories="、".join(company.get("categories") or []), types=types,
                status=STATUS_LABELS.get(company["status"], company["status"]))


def _channel(row):
    fields = {f["key"]: f["value"] for f in row.get("fields") or []}
    others = [f"{label} {fields[key]}" for key, label in (("phone", "电话"), ("whatsapp", "WhatsApp"),
                                                          ("extension", "分机")) if fields.get(key)]
    if row.get("email"):
        return "邮箱", row["email"], others
    if fields.get("phone"):
        return "电话", fields["phone"], [o for o in others if not o.startswith("电话 ")]
    if row.get("url"):
        kind = {"contact_form": "联系表单", "supplier_form": "供应商表单"}.get(row.get("kind"), "联系页面")
        return kind, row["url"], others
    return "无直接渠道（仅姓名/岗位）", "", others


def _review_label(row):
    if row.get("origin") == "manual":
        return "人工录入"
    if row.get("review_status") == "accept" or row.get("adoption_status") == "accepted":
        return "已人工确认"
    return "自动采集，未人工确认"


def _entry_sources(row, website):
    sources = sorted(row.get("sources") or [], key=lambda s: s.get("observed_at") or "", reverse=True)[:5]
    if row.get("origin") == "manual":
        labels = ["人工录入"]
    else:
        labels = list(dict.fromkeys(_source(s.get("source_url"), website) for s in sources))
    links = list(dict.fromkeys(s["source_url"] for s in sources if s.get("source_url")))
    observed = max((s.get("observed_at") or "" for s in sources), default=row.get("observed_at") or "")
    return "；".join(labels), "；".join(links), observed


def _entries_layout(data, format, hooks=None):
    hooks = hooks or _Progress()
    entries, company_rows, locations, sources = [], [], [], []
    total = len(data["companies"])
    for index, company in enumerate(data["companies"]):
        if index and index % PROGRESS_COMPANIES == 0:
            hooks.report("building", processed=index, total=total)
        columns = _company_columns(data, company)
        head = [company["name"], company["id"], columns["city"], columns["basis"], columns["market"],
                columns["categories"], columns["types"], columns["status"]]
        grouped = _grouped_contacts(data, company)
        for row in grouped:
            tier = contact_tier(row)[0]
            channel, value, others = _channel(row)
            label, links, observed = _entry_sources(row, company.get("website", ""))
            entries.append([safe(v) for v in head + [TIER_LABEL[tier], channel, value, "；".join(others),
                                                     row.get("name") or "", row.get("role") or "", label, links,
                                                     observed, _review_label(row), company.get("website") or ""]])
        if str(company.get("phone") or "").strip():
            listing = data["listings"].get(company["id"])
            phone_obs = data["phone_obs"].get(company["id"]) or {}
            source = "Google Maps" if listing else METHOD_LABELS.get(phone_obs.get("method"), "来源未记录")
            link = listing or phone_obs.get("source_url") or ""
            reviewed = "已人工确认" if "phone" in (company.get("overrides") or []) else "自动采集，未人工确认"
            entries.append([safe(v) for v in head + [TIER_LABEL["company_phone"], "电话", company["phone"], "", "", "",
                                                     source, link, phone_obs.get("observed_at") or "", reviewed,
                                                     company.get("website") or ""]])
        if format == "csv":
            continue
        latest = data["latest"].get(company["id"], {})
        company_rows.append([safe(v) for v in [
            company["name"], company["id"], company.get("address"), company.get("phone"), company.get("website"),
            columns["city"], columns["basis"], columns["market"], columns["categories"], columns["types"],
            columns["status"], company.get("best_tier_label"), company.get("contact_count"),
            (company.get("gap") or {}).get("reason", ""), latest.get("source_url"), latest.get("observed_at")]])
        for row in data["locations"].get(company["id"], []):
            locations.append([safe(v) for v in [company["name"], company["id"], row["address"], row["latitude"],
                                                 row["longitude"], row["source_url"], row["observed_at"]]])
        for row in data["observations"].get(company["id"], []):
            sources.append([safe(v) for v in [company["name"], company["id"], row["field"], row["value"],
                                               row["source_url"], row["evidence"], row["observed_at"],
                                               METHOD_LABELS.get(row["method"], row["method"]), row["query"]]])
    hooks.report("building", processed=total, total=total)
    return [("联系入口", ENTRY_HEADERS, entries), ("企业", ENTRY_COMPANY_HEADERS, company_rows),
            ("地点", ENTRY_LOCATION_HEADERS, locations), ("来源证据", ENTRY_SOURCE_HEADERS, sources)]
