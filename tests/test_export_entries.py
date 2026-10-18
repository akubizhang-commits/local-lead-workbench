"""P0-4: contact-entry export and legacy layout regression (design §5.4)."""
import csv
import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from leadworkbench.api import create_app
from leadworkbench.db import Store
from leadworkbench.export import ENTRY_HEADERS, export_data
from leadworkbench.sources.maps import _record
from legacy_db import FIXTURES, load_legacy


def rows(output):
    table = list(csv.reader(io.StringIO(output.decode("utf-8-sig"))))
    return [dict(zip(table[0], r)) for r in table[1:]], table[0]


def contact(name="", role="", email="", kind="general", role_status="unknown", page="https://alpha.example/team",
            evidence=None, url=""):
    return dict(name=name, role=role, email=email, url=url, kind=kind, role_status=role_status,
                association_status="same_block" if name else "company_only", source_url=page,
                evidence=evidence or f"{name} {role} {email} {url}".strip())


@pytest.fixture
def setup(tmp_path):
    store = Store(tmp_path)
    task = store.create_market(dict(country="US", state="California", city="San Jose", keywords=["roofing"],
                                    company_types=["wholesale"], auto_enrich=False))
    run = store.start_task(task["id"])

    def add(n, name, address, phone="", website="https://alpha.example", category="Roofing supply store", market=task):
        record = _record(dict(title=name, link=f"https://www.google.com/maps/search/?api=1&query_place_id=P{n}",
                              place_id=f"P{n}", address=address, phone=phone, web_site=website, category=category),
                         "2026-10-01T00:00:00+00:00")
        market_run = run if market is task else store.start_task(market["id"])
        return store.ingest_market(record, market_run["id"], "roofing", market["id"])["id"]
    return store, task, run, add


def test_one_row_per_entry_with_chinese_labels(setup):
    store, task, run, add = setup
    alpha = add(1, "Alpha Roofing", "100 Winchester Blvd, Campbell, CA 95008", phone="+1 408-555-0101")
    store.apply_enrichment(alpha, dict(status="completed", contacts=[
        contact("Bob Buyer", "Buyer", "bob@alpha.example", "procurement", "explicit",
                evidence="Bob Buyer | Buyer | Phone: +1 408 555 0177 | bob@alpha.example"),
        contact(email="info@alpha.example", page="https://alpha.example/contact"),
        contact(email="info@alpha.example", page="https://alpha.example/about"),
        contact(kind="contact_form", url="https://alpha.example/contact-us", page="https://alpha.example/contact-us",
                evidence="Contact us form")]))
    entries, headers = rows(export_data(store, "csv", "entries"))
    assert headers == ENTRY_HEADERS
    by_channel = {(e["渠道"], e["联系方式"]): e for e in entries}
    bob = by_channel[("邮箱", "bob@alpha.example")]
    assert (bob["入口等级"], bob["联系人"], bob["岗位"], bob["其他联系方式"]) == ("具名采购", "Bob Buyer", "Buyer", "电话 +14085550177")
    assert (bob["实际城市"], bob["城市依据"], bob["目标市场"]) == ("Campbell", "地址", task["name"])
    assert (bob["客户匹配状态"], bob["业务类型"], bob["来源类别"]) == ("候选企业", "批发商（待核实）", "Roofing supply store")
    assert bob["审核状态"] == "自动采集，未人工确认" and bob["来源"] == "企业官网"
    info = by_channel[("邮箱", "info@alpha.example")]
    assert info["入口等级"] == "通用邮箱"
    assert set(info["来源链接"].split("；")) == {"https://alpha.example/contact", "https://alpha.example/about"}
    assert by_channel[("联系表单", "https://alpha.example/contact-us")]["入口等级"] == "表单 / 联系页面 / 电话"
    phone = by_channel[("电话", "'+1 408-555-0101")]  # safe() keeps the formula guard (A1-F14)
    assert (phone["入口等级"], phone["来源"]) == ("仅企业电话", "Google Maps")
    assert phone["来源链接"].startswith("https://www.google.com/maps/")
    assert len(entries) == 4


def test_city_fallback_and_target_market_without_market_filter(setup):
    store, task, run, add = setup
    other = store.create_market(dict(country="US", state="California", city="San Jose", keywords=["metal"],
                                     company_types=[], auto_enrich=False))
    company = add(2, "Beta", "somewhere unparsable", phone="+1 408-555-0102")
    add(2, "Beta", "somewhere unparsable", phone="+1 408-555-0102", market=other)
    [entry], _ = rows(export_data(store, "csv", "entries"))
    assert (entry["实际城市"], entry["城市依据"]) == ("San Jose", "目标市场（地址未解析）")
    assert entry["目标市场"] == f"{task['name']}、{other['name']}"
    [entry], _ = rows(export_data(store, "csv", "entries", market_id=other["id"]))
    assert entry["目标市场"] == other["name"]
    assert company


def test_exclusions_refusals_and_rejections(setup):
    store, task, run, add = setup
    kept = add(3, "Kept", "1 A St, San Jose, CA 95112", phone="+1 408-555-0103")
    refused = add(4, "Refused", "2 B St, San Jose, CA 95112", phone="+1 408-555-0104")
    store.update_company(refused, store.company(refused)["revision"], {}, status="do_not_contact")
    hidden = add(5, "Hidden Contractor", "3 C St, San Jose, CA 95112", phone="+1 408-555-0105",
                 category="Roofing contractor")
    store.create_exclusion_rule("global", "category", "Roofing contractor")
    store.apply_enrichment(kept, dict(status="completed", contacts=[contact(email="out@alpha.example")]))
    with store.connect() as con:
        con.execute("INSERT INTO contact_decisions SELECT id,'excluded','2026' FROM contacts WHERE email='out@alpha.example'")
    store.apply_enrichment(kept, dict(status="completed", contacts=[contact(email="jobs@alpha.example")]))
    entries, _ = rows(export_data(store, "csv", "entries"))
    assert [e["企业名称"] for e in entries] == ["Kept"]
    assert [e["入口等级"] for e in entries] == ["仅企业电话"]
    assert hidden


def test_xlsx_has_four_sheets_each_starting_with_company_name(setup):
    store, task, run, add = setup
    alpha = add(1, "=cmd|' /C calc'!A0", "100 Main St, San Jose, CA 95112")
    store.apply_enrichment(alpha, dict(status="completed", contacts=[contact(email="info@alpha.example")]))
    book = load_workbook(io.BytesIO(export_data(store, "xlsx", "entries")))
    assert book.sheetnames == ["联系入口", "企业", "地点", "来源证据"]
    for sheet in book:
        values = list(sheet.values)
        assert values[0][0] == "企业名称" and values[1][0] == "'=cmd|' /C calc'!A0"
        assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref.startswith("A1:")


def test_manual_contact_exported_as_manual(setup):
    store, task, run, add = setup
    company = add(1, "Manual", "1 A St, San Jose, CA 95112")
    store.add_manual_contact(company, store.company(company)["revision"], name="Mia Manual", email="mia@alpha.example",
                             procurement=True, note="trade show")
    [entry], _ = rows(export_data(store, "csv", "entries"))
    assert (entry["来源"], entry["审核状态"], entry["入口等级"]) == ("人工录入", "人工录入", "具名采购")


def test_legacy_layout_unchanged_on_export_source_fixture(tmp_path):
    load_legacy(tmp_path, "legacy_export_source.sql")
    book = load_workbook(io.BytesIO(export_data(Store(tmp_path), "xlsx")))
    baseline = json.loads((FIXTURES / "legacy_export_9423c53.json").read_text(encoding="utf-8"))
    for name in ("Contacts", "Locations", "Sources"):
        got = [[c if c is not None else "" for c in r] for r in book[name].iter_rows(values_only=True)]
        assert got[0][-1] == "company_name"
        assert [r[:-1] for r in got] == baseline[name], name
    assert book.sheetnames == ["Companies", "Contacts", "Locations", "Sources"]


def test_legacy_layout_on_full_fixture_only_adds_repaired_contacts(tmp_path):
    load_legacy(tmp_path)
    store = Store(tmp_path)
    with store.connect() as con:
        repaired = {r[0] for r in con.execute("SELECT c.email FROM contact_change_proposals p "
                                               "JOIN contacts c ON c.id=p.contact_id WHERE p.origin='repair'")}
    assert repaired == {"bob@alpha.example", "dana@gamma.example"}
    book = load_workbook(io.BytesIO(export_data(store, "xlsx")))
    baseline = json.loads((FIXTURES / "legacy_v2_export_9423c53.json").read_text(encoding="utf-8"))
    contacts = [[c if c is not None else "" for c in r][:-1] for r in book["Contacts"].iter_rows(values_only=True)]
    assert sorted(map(tuple, [r for r in contacts if r[3] not in repaired])) == sorted(map(tuple, baseline["Contacts"]))
    assert {r[1] for r in contacts if r[3] in repaired} == {"Bob Buyer", "Dana Smith"}
    sources = [[c if c is not None else "" for c in r][:-1] for r in book["Sources"].iter_rows(values_only=True)]
    assert sorted(map(tuple, [r for r in sources if not (r[1] == "contact" and r[2] in repaired)])) == \
        sorted(map(tuple, baseline["Sources"]))


def test_api_layout_and_filename(tmp_path):
    app = create_app(tmp_path, allowed_hosts={"testserver"})
    with TestClient(app) as client:
        client.headers["X-Local-Token"] = client.get("/api/session").json()["token"]
        response = client.get("/api/exports", params={"format": "csv", "layout": "entries", "best_tier": "none"})
        assert response.status_code == 200
        assert 'filename="contact-entries.csv"' in response.headers["content-disposition"]
        assert response.text.lstrip("\ufeff").startswith("企业名称,")
        legacy = client.get("/api/exports", params={"format": "csv"})
        assert 'filename="leads.csv"' in legacy.headers["content-disposition"]
        assert client.get("/api/exports", params={"layout": "other"}).status_code == 422
