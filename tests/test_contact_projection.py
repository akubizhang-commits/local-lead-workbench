from leadworkbench.contact_projection import group_contact_records
from leadworkbench.db import Store


def contact(**values):
    return dict(id='a', company_id='co', name='Alice Example', role='Buyer', email='alice@example.org',
                url='', kind='procurement', role_status='explicit', association_status='same_block',
                source_url='https://example.org/team', evidence='Alice Example Buyer',
                observed_at='2026-09-22', review_status='auto_recorded', **values)


def test_same_contact_multiple_sources_one_card_preserves_evidence():
    a = contact()
    b = {**a, 'id':'b', 'source_url':'https://example.org/contact', 'evidence':'Alice Example Buyer Phone: 2134543667'}
    result = group_contact_records([a,b], 'https://example.org')
    assert len(result) == 1
    assert len(result[0]['sources']) == 2
    assert next(f for f in result[0]['fields'] if f['key']=='phone')['source_url'] == b['source_url']


def test_conflicting_people_roles_and_channels_do_not_merge():
    a = contact()
    variants = [{**a, 'id':'b', 'name':'Bob Example'}, {**a, 'id':'c', 'role':'Sales Manager'},
                {**a,'id':'d','email':'second@example.org'}]
    assert len(group_contact_records([a,*variants])) == 4


def test_rejected_and_invalid_are_not_contacts():
    a = contact()
    assert group_contact_records([{**a,'review_status':s} for s in ['reject','invalid','superseded']]) == []


def test_contact_cleanup_retains_source_and_history_and_is_idempotent(tmp_path):
    s=Store(tmp_path); co=s.manual_company('Fixture','https://example.org')
    with s.connect() as con:
        con.execute('INSERT INTO contacts VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
                    ('bad',co['id'],'Quality Control','Plan smaller production runs with realistic sourcing options.',
                     '', '', 'procurement','explicit','same_block','https://example.org/services',
                     'Quality Control Plan smaller production runs with realistic sourcing options.',
                     '2026-09-22','auto_recorded','bad'))
    assert s.reclassify_contacts()['archived'] == 1
    assert s.reclassify_contacts()['archived'] == 0
    assert s.detail(co['id'])['contacts'] == []
    with s.connect() as con:
        assert con.execute("SELECT evidence FROM contacts WHERE id='bad'").fetchone()[0]
    assert any(c['reason'].startswith('联系人识别修正') for c in s.changes(co['id']))


def test_grouped_count_and_export_match_detail(tmp_path):
    import io
    from openpyxl import load_workbook
    from leadworkbench.export import export_data
    s=Store(tmp_path); co=s.manual_company('Fixture','https://example.org')
    a=contact(); b={**a,'source_url':'https://example.org/contact'}
    s.apply_enrichment(co['id'], {'status':'completed','contacts':[a,b]})
    assert s.company(co['id'])['contact_count']==1
    assert s.company_page()['items'][0]['contact_count']==1
    assert len(s.detail(co['id'])['contacts'])==1
    book=load_workbook(io.BytesIO(export_data(s,'xlsx')))
    assert book['Contacts'].max_row==2
    assert sum(r[1]=='contact' for r in book['Sources'].iter_rows(min_row=2,values_only=True))==2


def test_rejected_contact_is_not_revived_by_new_evidence(tmp_path):
    s=Store(tmp_path); co=s.manual_company('Fixture','https://example.org')
    a=contact()
    s.apply_enrichment(co['id'],{'status':'completed','contacts':[a]})
    with s.connect() as con:
        con.execute("UPDATE contacts SET review_status='reject'")
    s.apply_enrichment(co['id'],{'status':'completed','contacts':[{**a,'role':'Sales Manager'}]})
    # Excluded contacts remain inspectable, but are never usable/exported.
    assert s.detail(co['id'])['contacts'][0]['adoption_status']=='excluded'
    assert s.company(co['id'])['contact_count']==0
    assert s.company_page(contact='procurement')['total']==0


def test_incorrect_person_preserves_actual_public_mailbox(tmp_path):
    s=Store(tmp_path); co=s.manual_company('Fixture','https://example.org')
    a={**contact(), 'name':'Quality Control', 'role':'Plan smaller production runs with realistic sourcing options.',
       'evidence':'Quality Control Email: alice@example.org'}
    s.apply_enrichment(co['id'],{'status':'completed','contacts':[a]})
    contacts=s.detail(co['id'])['contacts']
    assert len(contacts)==1
    assert contacts[0]['email']=='alice@example.org'
    assert contacts[0]['name']=='' and contacts[0]['kind']=='general'
    assert contacts[0]['association_status']=='company_only'


def test_merge_versions_attaches_only_the_same_identity():
    from leadworkbench.contact_projection import merge_versions
    bob = dict(id='k', name='Bob Buyer', role='Buyer', email='team@x.example', kind='procurement',
               role_status='explicit', association_status='same_block', source_url='https://x.example/b',
               evidence='Bob', observed_at='2')
    alex = {**bob, 'name': 'Alex Seller', 'role': 'Sales Manager', 'kind': 'general',
            'evidence': 'Alex | Phone: +1 408 555 0188', 'source_url': 'https://x.example/a', 'observed_at': '1'}
    again = {**bob, 'source_url': 'https://x.example/c', 'evidence': 'Bob again', 'observed_at': '3'}
    records = merge_versions([bob], {'k': [alex, again]})
    assert [r['source_url'] for r in records] == ['https://x.example/b', 'https://x.example/c']
    assert all('0188' not in r['evidence'] for r in records)
