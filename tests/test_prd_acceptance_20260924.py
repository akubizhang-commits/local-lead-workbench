"""PRD acceptance regressions: temporary databases, ASGI, synthetic source data.

These tests cover implemented behavior, never delivery or real source availability.
"""
import csv
import io
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from leadworkbench.api import create_app
from leadworkbench.db import Store
from leadworkbench.export import export_data


@pytest.fixture
def store(tmp_path):
    return Store(tmp_path)


@pytest.fixture
def client(tmp_path):
    app = create_app(tmp_path, allowed_hosts={"testserver"})
    # No context manager/lifespan: never open or close a browser session.
    client = TestClient(app, raise_server_exceptions=False)
    try:
        client.headers["X-Local-Token"] = client.get("/api/session").json()["token"]
        yield client, app.state.store
    finally:
        client.close()


def market(store, city="Los Angeles"):
    return store.create_market(dict(
        name="Acceptance fixture", country="US", state="California", city=city,
        keywords=["apparel"], company_types=["wholesale"], auto_enrich=False,
    ))


def listing(key="one", **extra):
    return dict(
        source_key=f"maps:{key}", source_url=f"https://maps.example/{key}",
        name=f"Fixture {key}", address=f"{key} Test Street",
        website=f"https://{key}.example", phone="",
        observed_at="2026-09-20T00:00:00+00:00", **extra,
    )


def contact(email="alex@one.example", **extra):
    return {
        "name": "Alex Chen", "role": "Buyer", "email": email, "kind": "procurement",
        "role_status": "explicit", "association_status": "same_block",
        "source_url": "https://one.example/team",
        "evidence": f"Alex Chen Buyer {email}",
        "observed_at": "2026-09-20T00:00:00+00:00", **extra,
    }


def apply(store, cid, **payload):
    store.apply_enrichment(cid, {"status": "completed", **payload})


def csv_rows(data):
    return list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))


def sheet_rows(book, name):
    rows = list(book[name].values)
    return [dict(zip(rows[0], row)) for row in rows[1:]]


def test_prd03_replay_empty_fill_manual_conflict_and_market_membership(store):
    first_market = market(store)
    run = store.start_task(first_market["id"])
    record = listing()
    company = store.ingest_market(record, run["id"], "fixture query", first_market["id"])
    cid = company["id"]
    assert store.ingest_market(record, run["id"], "fixture query", first_market["id"])["id"] == cid
    filled = {**record, "phone": "+1 212 555 0101", "observed_at": "2026-09-21T00:00:00+00:00"}
    store.ingest_market(filled, run["id"], "fixture query", first_market["id"])
    assert store.company(cid)["phone"] == filled["phone"]
    assert not store.reviews()
    current = store.company(cid)
    store.update_company(cid, current["revision"], {"phone": "+1 212 555 0199"}, notes="Human decision")
    store.ingest_market({**filled, "phone": "+1 212 555 0188",
                         "observed_at": "2026-09-22T00:00:00+00:00"},
                        run["id"], "fixture query", first_market["id"])
    review = next(r for r in store.reviews() if r["field"] == "phone")
    assert review["old_value"] == "+1 212 555 0199"
    assert review["new_value"] == "+1 212 555 0188"
    assert store.company(cid)["notes"] == "Human decision"
    assert store.company(cid)["phone"] == review["old_value"]
    second_market = market(store, city="San Francisco")
    second_run = store.start_task(second_market["id"])
    assert store.ingest_market(filled, second_run["id"], "other market", second_market["id"])["id"] == cid
    assert len(store.companies()) == 1
    assert store.company_page(market_id=second_market["id"], segment="new")["total"] == 1
    assert store.company_page(market_id=second_market["id"], segment="existing")["total"] == 0
    assert any(o["value"] == review["new_value"] and o["source_url"] == record["source_url"]
               for o in store.detail(cid)["observations"])


@pytest.mark.parametrize("status", ["failed", "blocked", "auth_required"])
def test_prd07_failed_website_payload_cannot_replace_successful_data(store, status):
    cid = store.manual_company("Fixture", "https://one.example")["id"]
    apply(store, cid, contacts=[contact()], observations=[{
        "field": "phone", "value": "12125550101", "source_url": "https://one.example/contact",
        "evidence": "Office phone", "observed_at": "2026-09-20T00:00:00+00:00",
    }], social_links=["https://x.com/fixture"])
    before = store.detail(cid)
    apply(store, cid, status=status, contacts=[contact(email="bad@one.example")],
          observations=[{"field": "phone", "value": "", "source_url": "https://one.example",
                         "evidence": "Failed page"}], social_links=[])
    after = store.detail(cid)
    assert after == before
    assert not store.reviews()


def test_prd07_empty_contact_fill_conflict_and_export_adoption(store):
    cid = store.manual_company("Fixture", "https://one.example")["id"]
    original = contact(role="", role_status="")
    apply(store, cid, contacts=[original, original])
    apply(store, cid, contacts=[contact()])
    assert not store.reviews()
    assert len(store.detail(cid)["contacts"]) == 1
    assert store.company_page(contact="procurement")["total"] == 1
    apply(store, cid, contacts=[contact(role="Sales Manager", kind="email", role_status="unknown")])
    review = store.reviews()[0]
    assert review["previous_contact"]["role"] == "Buyer"
    assert review["contact"]["role"] == "Sales Manager"
    book = load_workbook(io.BytesIO(export_data(store, "xlsx")))
    assert sheet_rows(book, "Contacts") == []
    store.review_decision(review["id"], "reject")
    apply(store, cid, contacts=[contact(evidence="New page with the same public mailbox")])
    assert store.detail(cid)["contacts"][0]["adoption_status"] == "excluded"
    assert sheet_rows(load_workbook(io.BytesIO(export_data(store, "xlsx"))), "Contacts") == []


@pytest.mark.parametrize("source_status,item_status", [
    ("preview_read", "succeeded"), ("auth_required", "blocked"),
    ("challenge", "blocked"), ("unreadable", "blocked"),
    ("failed", "failed"), ("no_results", "no_results"),
])
def test_prd04_source_status_retains_reason_and_does_not_confirm_identity(store, source_status, item_status):
    cid = store.manual_company("Fixture", "https://one.example")["id"]
    store.add_social(cid, "https://x.com/fixture", "Synthetic source")
    run = store.create_social_run([cid], read_only=True)
    assert run["summary"]["pending"] == 1 and not store.social_pages(cid)
    item = store.claim_social_item()
    assert store.finish_social_item(item["id"], {"status": source_status, "message": "Fixture reason", "fields": []})
    result = store.social_run(run["id"])
    assert result["summary"][item_status] == 1
    assert result["items"][0]["result"]["status"] == source_status
    assert result["items"][0]["result"]["message"] == "Fixture reason"
    detail = store.social_detail(cid)
    assert detail["accounts"][0]["status"] == "candidate"
    assert detail["capabilities"]["auto_send"] is False
    assert store.detail(cid)["contacts"] == []


def test_prd10_refusal_rejects_late_social_result_and_disables_due_schedule(store):
    cid = store.manual_company("Fixture", "https://one.example")["id"]
    entry = store.add_social(cid, "https://x.com/fixture", "Synthetic source")
    store.decide_social(cid, entry["id"], "confirmed", "brand", revision=1)
    store.record_social_outreach(cid, entry["id"], "draft", "Existing manual draft")
    run = store.create_social_run([cid], read_only=True)
    item = store.claim_social_item()
    schedule = store.create_schedule(company_id=cid, kind="refresh", timezone="UTC")
    current = store.company(cid)
    store.update_company(cid, current["revision"], {}, status="do_not_contact")
    assert not store.finish_social_item(item["id"], {
        "status": "preview_read", "fields": [{"key": "email", "value": "late@one.example"}],
        "posts": [{"url": "https://x.com/fixture/status/1", "text": "Late result"}],
    })
    assert store.social_run(run["id"])["items"][0]["state"] == "cancelled"
    assert not store.social_pages(cid)
    assert len(store.social_detail(cid)["outreach"]) == 1
    with pytest.raises(ValueError, match="拒绝联系"):
        store.record_social_outreach(cid, entry["id"], "contacted", "Must not be saved")
    with store.connect() as con:
        con.execute("UPDATE schedules SET next_due_at='2020-01-01T00:00:00+00:00' WHERE id=?", (schedule["id"],))
    store.dispatch_due()
    store.dispatch_due()
    assert store.schedules()[0]["enabled"] is False
    assert store.runs() == []
    with store.connect() as con:
        assert con.execute("SELECT count(*) FROM occurrences").fetchone()[0] == 0


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_prd11_api_export_matches_combined_filter_and_excludes_refusal(client, format):
    client, store = client
    m = market(store)
    run = store.start_task(m["id"])
    selected = store.ingest_market(listing(), run["id"], "fixture", m["id"])
    refused = store.ingest_market(listing("refused"), run["id"], "fixture", m["id"])
    outside = store.manual_company("Fixture outside", "https://outside.example")
    for company in (selected, refused, outside):
        apply(store, company["id"], contacts=[contact(email=f"alex@{company['id']}.example")])
    store.update_company(refused["id"], store.company(refused["id"])["revision"], {}, status="do_not_contact")
    params = dict(run_id=run["id"], segment="new", contact="procurement", q="Fixture", sort="procurement")
    page = client.get("/api/company-page", params={**params, "page_size": 1}).json()
    assert page["total"] == 2 and len(page["items"]) == 1
    response = client.get("/api/exports", params={**params, "format": format})
    assert response.status_code == 200
    assert f'filename="leads.{format}"' in response.headers["content-disposition"]
    if format == "csv":
        rows = csv_rows(response.content)
    else:
        book = load_workbook(io.BytesIO(response.content))
        assert book.sheetnames == ["Companies", "Contacts", "Locations", "Sources"]
        rows = sheet_rows(book, "Companies")
        for name in ("Contacts", "Locations", "Sources"):
            related = sheet_rows(book, name)
            assert related
            assert {r["company_id"] for r in related} == {selected["id"]}
        assert sheet_rows(book, "Contacts")[0]["email"].endswith(".example")
    assert [row["id"] for row in rows] == [selected["id"]]
    assert rows[0]["market_segment"] == "new"
    assert rows[0]["source_url"] and rows[0]["observed_at"]
    assert client.get("/api/exports", params={**params, "q": "not present"}).status_code == 200


def test_prd11_backup_restore_preserves_provenance_and_human_decisions(client):
    client, store = client
    cid = store.manual_company("Fixture", "https://one.example")["id"]
    apply(store, cid, contacts=[contact()])
    apply(store, cid, contacts=[contact(role="Purchasing Manager")])
    store.review_decision(store.reviews()[0]["id"], "reject")
    entry = store.add_social(cid, "https://x.com/fixture", "Manual source evidence")
    store.decide_social(cid, entry["id"], "confirmed", "brand", revision=1)
    store.record_social_outreach(cid, entry["id"], "draft", "Persisted manual draft")
    store.update_company(cid, store.company(cid)["revision"], {"phone": "12125550101"}, notes="Human note")
    before = store.detail(cid)
    social_before = store.social_detail(cid)
    backup = client.post("/api/backups").json()
    store.update_company(cid, store.company(cid)["revision"], {"phone": "12125550999"}, notes="Later")
    store.decide_social(cid, entry["id"], "excluded", "third_party", revision=2)
    assert client.post("/api/backups/restore", json={"name": backup["name"]}).status_code == 200
    assert store.detail(cid) == before
    assert store.social_detail(cid) == social_before
    assert sheet_rows(load_workbook(io.BytesIO(export_data(store, "xlsx"))), "Contacts") == []
    with store.connect() as con:
        assert con.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        assert con.execute("PRAGMA foreign_key_check").fetchall() == []


@pytest.mark.parametrize("headers,expected", [
    ({"Host": "evil.example"}, 403),
    ({"Origin": "https://evil.example"}, 403),
    ({"Sec-Fetch-Site": "cross-site"}, 403),
    ({"X-Local-Token": "wrong"}, 401),
    ({"Content-Length": "1000001"}, 413),
    ({"Content-Length": "nonnumeric"}, 413),
])
def test_api_boundary_rejects_write_before_mutation(client, headers, expected):
    client, store = client
    response = client.post("/api/companies", json={"name": "Must not persist", "website": "https://one.example"},
                           headers=headers)
    assert response.status_code == expected
    assert store.companies() == []


@pytest.mark.parametrize("name", ["../outside.sqlite3", "/tmp/outside.sqlite3", "missing.sqlite3", "invalid.txt"])
def test_api_restore_invalid_name_leaves_data_intact(client, name):
    client, store = client
    cid = store.manual_company("Fixture", "https://one.example")["id"]
    before = store.detail(cid)
    assert client.post("/api/backups/restore", json={"name": name}).status_code == 400
    assert store.detail(cid) == before


def test_api_restore_symlink_and_active_refresh_are_rejected(client):
    client, store = client
    cid = store.manual_company("Fixture", "https://one.example")["id"]
    backup = client.post("/api/backups").json()
    (store.directory / "backups" / "alias.sqlite3").symlink_to(backup["name"])
    assert client.post("/api/backups/restore", json={"name": "alias.sqlite3"}).status_code == 400
    run = client.post(f"/api/companies/{cid}/refresh").json()
    assert client.post("/api/backups/restore", json={"name": backup["name"]}).status_code == 409
    assert store.run(run["id"])["status"] == "queued"
    assert client.post(f"/api/runs/{run['id']}/cancel").status_code == 200
    assert client.post("/api/backups/restore", json={"name": backup["name"]}).status_code == 200


@pytest.mark.parametrize("extra,expected", [
    ({"company_id": "missing", "task_id": "missing", "kind": "refresh"}, 400),
    ({"kind": "refresh"}, 400),
    ({"company_id": "missing", "kind": "discovery"}, 400),
    ({"company_id": "missing", "kind": "refresh", "timezone": "No/SuchZone"}, 422),
    ({"company_id": "missing", "kind": "refresh", "interval_months": 0}, 422),
])
def test_api_schedule_invalid_definition_never_creates_schedule(client, extra, expected):
    client, store = client
    assert client.post("/api/schedules", json=extra).status_code == expected
    assert store.schedules() == []
