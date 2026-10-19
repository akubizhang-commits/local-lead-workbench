"""P1-2: background export jobs in a child process (design §8.9)."""
import asyncio
from datetime import datetime, timedelta, timezone
import io
import json
import os
import signal
import sqlite3
import time

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from leadworkbench import export_jobs
from leadworkbench.api import create_app
from leadworkbench.db import Store
from leadworkbench.export import _collect, export_data
from perf_fixtures import build

TERMINAL = {"succeeded", "failed", "cancelled"}


@pytest.fixture
def small(tmp_path):
    build(tmp_path, companies=60, contacts_per_company=2, markets=1)
    return tmp_path


@pytest.fixture
def large(tmp_path):
    build(tmp_path, companies=2000, contacts_per_company=2, markets=1)
    return tmp_path


def client_for(directory):
    app = create_app(directory, allowed_hosts={"testserver"})
    client = TestClient(app, raise_server_exceptions=False)
    client.__enter__()
    client.headers["X-Local-Token"] = client.get("/api/session").json()["token"]
    return client, app


@pytest.fixture
def api(small):
    client, app = client_for(small)
    try:
        yield client, app, small
    finally:
        client.__exit__(None, None, None)


@pytest.fixture
def api_large(large):
    client, app = client_for(large)
    try:
        yield client, app, large
    finally:
        client.__exit__(None, None, None)


def wait_job(client, job_id, predicate=lambda job: job["status"] in TERMINAL, timeout=60):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        job = client.get(f"/api/export-jobs/{job_id}").json()
        if predicate(job):
            return job
        time.sleep(0.1)
    raise AssertionError(f"job {job_id} did not reach the expected state: {job}")


def create(client, **body):
    response = client.post("/api/export-jobs", json={"format": "csv", "layout": "entries", **body})
    assert response.status_code == 202, response.text
    return response.json()


def no_child_processes():
    import subprocess
    out = subprocess.run(["ps", "-A", "-o", "command="], capture_output=True, text=True).stdout
    return "leadworkbench.export_job" not in out


def sheets(data):
    book = load_workbook(io.BytesIO(data), read_only=True)
    return {ws.title: [list(row) for row in ws.iter_rows(values_only=True)] for ws in book.worksheets}


@pytest.mark.parametrize("format,layout", [("csv", "entries"), ("csv", "companies"), ("xlsx", "entries"),
                                           ("xlsx", "companies")])
def test_1_job_file_equals_sync_export(api, format, layout):
    client, app, directory = api
    job = create(client, format=format, layout=layout)
    assert job["status"] == "queued" and job["label"].endswith(format.upper())
    done = wait_job(client, job["id"])
    assert done["status"] == "succeeded", done
    assert done["total"] == done["processed"] > 0 and done["size_bytes"] > 0
    assert done["expires_at"] > done["finished_at"]
    response = client.get(f"/api/export-jobs/{job['id']}/file")
    assert response.status_code == 200
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    expected = export_data(app.state.store, format, layout)
    if format == "csv":
        assert response.content == expected
    else:
        assert sheets(response.content) == sheets(expected)
    final = export_jobs.file_path(directory, done)
    assert oct(final.stat().st_mode & 0o777) == "0o600"
    assert oct(export_jobs.exports_dir(directory).stat().st_mode & 0o777) == "0o700"
    assert oct(export_jobs.job_path(directory, job["id"]).stat().st_mode & 0o777) == "0o600"


def test_2_second_job_is_busy(api_large, monkeypatch):
    client, app, directory = api_large
    monkeypatch.setenv("LEAD_TEST_FAULT", "export_slow")
    job = create(client, format="xlsx")
    response = client.post("/api/export-jobs", json={"format": "csv", "layout": "entries"})
    assert response.status_code == 409
    assert response.json()["code"] == "export_busy" and response.json()["job_id"] == job["id"]
    client.post(f"/api/export-jobs/{job['id']}/cancel")
    wait_job(client, job["id"])


def test_3_cancel_during_building(api_large, monkeypatch):
    client, app, directory = api_large
    monkeypatch.setenv("LEAD_TEST_FAULT", "export_slow")
    job = create(client, format="xlsx")
    wait_job(client, job["id"], lambda j: j["phase"] == "building")
    assert client.post(f"/api/export-jobs/{job['id']}/cancel").status_code == 200
    done = wait_job(client, job["id"])
    assert done["status"] == "cancelled" and done["error"] is None
    assert not export_jobs.part_path(directory, job["id"]).exists()
    assert client.post(f"/api/export-jobs/{job['id']}/cancel").status_code == 409


def test_4_restart_marks_stale_job_interrupted(small):
    export_jobs.ensure_dirs(small)
    stale = (datetime.now(timezone.utc) - timedelta(seconds=60)).isoformat()
    job_id = "a" * 32
    export_jobs.write_job(small, dict(id=job_id, status="running", phase="building", format="csv", layout="entries",
                                      filters={}, created_at=stale, updated_at=stale))
    export_jobs.part_path(small, job_id).write_text("partial")
    app = create_app(small, allowed_hosts={"testserver"})
    job = export_jobs.read_job(export_jobs.job_path(small, job_id))
    assert job["status"] == "failed" and job["error"]["code"] == "interrupted"
    assert "本地服务已重启" in job["error"]["message"]
    assert not export_jobs.part_path(small, job_id).exists()
    assert app.state.export_jobs.list()[0]["id"] == job_id


def test_4b_killed_child_is_interrupted_without_restart(api_large, monkeypatch):
    client, app, directory = api_large
    monkeypatch.setenv("LEAD_TEST_FAULT", "export_slow")
    job = create(client, format="xlsx")
    running = wait_job(client, job["id"], lambda j: j["phase"] == "building")
    os.kill(running["pid"], signal.SIGKILL)
    started = time.monotonic()
    done = wait_job(client, job["id"], timeout=3)
    assert time.monotonic() - started < 3
    assert done["status"] == "failed" and done["error"]["code"] == "interrupted"
    assert "导出进程意外退出" in done["error"]["message"]
    assert not export_jobs.part_path(directory, job["id"]).exists()
    monkeypatch.delenv("LEAD_TEST_FAULT")
    second = create(client)
    wait_job(client, second["id"])
    restore = client.post("/api/backups/restore", json={"name": "backup-missing.sqlite3"})
    assert restore.json()["code"] != "export_running"


async def test_4c_concurrent_creates_make_one_job(small, monkeypatch):
    calls = []
    class FakeProcess:
        pid = 4242
        def __init__(self, *args, **kwargs):
            time.sleep(0.05)
            calls.append(args)
        def poll(self):
            return None
        def terminate(self):
            pass
        def wait(self, timeout=None):
            return 0
    monkeypatch.setattr(export_jobs.subprocess, "Popen", FakeProcess)
    app = create_app(small, allowed_hosts={"testserver"})
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://testserver") as http:
        http.headers["X-Local-Token"] = (await http.get("/api/session")).json()["token"]
        responses = await asyncio.gather(*(http.post("/api/export-jobs", json={"format": "csv"}) for _ in range(20)))
    created = [r for r in responses if r.status_code == 202]
    busy = [r for r in responses if r.status_code == 409]
    assert len(created) == 1 and len(busy) == 19 and len(calls) == 1
    assert {r.json()["job_id"] for r in busy} == {created[0].json()["id"]}
    assert len(list(export_jobs.jobs_dir(small).glob("*.json"))) == 1


def test_4d_heartbeat_during_blocking_save(api, monkeypatch):
    client, app, directory = api
    monkeypatch.setenv("LEAD_TEST_FAULT", "export_save_slow")
    job = create(client, format="xlsx")
    wait_job(client, job["id"], lambda j: j["phase"] == "writing")
    stamps = set()
    deadline = time.monotonic() + 4.5
    while time.monotonic() < deadline:
        current = client.get(f"/api/export-jobs/{job['id']}").json()
        if current["status"] in TERMINAL:
            break
        stamps.add(current["updated_at"])
        time.sleep(0.2)
    assert len(stamps) >= 3, stamps  # at least two heartbeat refreshes
    assert wait_job(client, job["id"])["status"] == "succeeded"


def test_4e_dirty_facets_still_match(api):
    client, app, directory = api
    store = app.state.store
    with store.connect() as con:
        con.execute("UPDATE contacts SET role='Purchasing Manager' WHERE rowid % 3 = 0")
        assert con.execute("SELECT count(*) FROM contact_facets_dirty").fetchone()[0] > 0
    job = create(client)
    done = wait_job(client, job["id"])
    assert done["status"] == "succeeded"
    assert client.get(f"/api/export-jobs/{job['id']}/file").content == export_data(store, "csv", "entries")


def test_4f_ignored_cancel_gets_sigterm(api_large, monkeypatch):
    client, app, directory = api_large
    monkeypatch.setenv("LEAD_TEST_FAULT", "export_ignore_cancel")
    monkeypatch.setattr(export_jobs, "CANCEL_TERM_SECONDS", 1)
    job = create(client, format="xlsx")
    wait_job(client, job["id"], lambda j: j["phase"] == "building")
    client.post(f"/api/export-jobs/{job['id']}/cancel")
    done = wait_job(client, job["id"], timeout=10)
    assert done["status"] == "cancelled"
    assert not export_jobs.part_path(directory, job["id"]).exists()


def test_5_restore_refuses_while_exporting(api_large, monkeypatch):
    client, app, directory = api_large
    monkeypatch.setenv("LEAD_TEST_FAULT", "export_slow")
    job = create(client, format="xlsx")
    response = client.post("/api/backups/restore", json={"name": "backup-x.sqlite3"})
    assert response.status_code == 409 and response.json()["code"] == "export_running"
    client.post(f"/api/export-jobs/{job['id']}/cancel")
    wait_job(client, job["id"])


def test_6_download_states(api_large, monkeypatch):
    client, app, directory = api_large
    assert client.get("/api/export-jobs/not-a-job/file").status_code == 404
    assert client.get("/api/export-jobs/..%2F..%2Fetc/file").status_code == 404
    monkeypatch.setenv("LEAD_TEST_FAULT", "export_slow")
    job = create(client, format="xlsx")
    pending = client.get(f"/api/export-jobs/{job['id']}/file")
    assert pending.status_code == 409 and pending.json()["code"] == "conflict"
    client.post(f"/api/export-jobs/{job['id']}/cancel")
    wait_job(client, job["id"])
    monkeypatch.delenv("LEAD_TEST_FAULT")
    done = wait_job(client, create(client)["id"])
    export_jobs.file_path(directory, done).unlink()
    gone = client.get(f"/api/export-jobs/{done['id']}/file")
    assert gone.status_code == 410 and gone.json()["code"] == "export_expired"
    assert client.get(f"/api/export-jobs/{'b' * 32}/file").status_code == 410
    assert client.get(f"/api/export-jobs/{'b' * 32}").status_code == 404


def test_7_disk_full_fails_cleanly(api, monkeypatch):
    client, app, directory = api
    monkeypatch.setenv("LEAD_TEST_FAULT", "export_disk_full")
    done = wait_job(client, create(client)["id"])
    assert done["status"] == "failed" and done["error"]["code"] == "disk_full"
    assert "磁盘空间不足" in done["error"]["message"]
    assert not export_jobs.part_path(directory, done["id"]).exists()


def test_8_sync_export_is_deprecated_but_works(api):
    client, app, directory = api
    response = client.get("/api/exports?format=csv&layout=entries")
    assert response.status_code == 200 and response.headers["Deprecation"] == "true"
    assert "successor-version" in response.headers["Link"]


def test_invalid_filters_fail_before_a_job_starts(api):
    client, app, directory = api
    assert client.post("/api/export-jobs", json={"format": "csv", "run_id": "missing"}).status_code == 404
    assert client.post("/api/export-jobs", json={"format": "csv", "market_id": "missing"}).json()["code"] == "not_found"
    assert client.post("/api/export-jobs", json={"format": "pdf"}).status_code == 422
    assert client.post("/api/export-jobs", json={"format": "csv", "extra": 1}).status_code == 422
    assert not list(export_jobs.jobs_dir(directory).glob("*.json")) if export_jobs.jobs_dir(directory).exists() else True


def test_child_ignores_sigint(api_large, monkeypatch):
    client, app, directory = api_large
    monkeypatch.setenv("LEAD_TEST_FAULT", "export_slow")
    job = create(client, format="csv")
    running = wait_job(client, job["id"], lambda j: j["phase"] == "building")
    os.kill(running["pid"], signal.SIGINT)
    time.sleep(0.5)
    assert client.get(f"/api/export-jobs/{job['id']}").json()["status"] == "running"
    assert wait_job(client, job["id"], timeout=120)["status"] == "succeeded"


def test_read_snapshot_holds_no_write_lock(small):
    store = Store(small)
    with store.connect() as reader:
        reader.execute("BEGIN")
        data = _collect(store, reader, {}, observations=False)
        assert reader.execute("SELECT count(*) FROM temp.export_ids").fetchone()[0] == len(data["companies"])
        before = reader.execute("SELECT name FROM companies WHERE id='c000001'").fetchone()[0]
        writer = sqlite3.connect(store.path, timeout=15)
        start = time.monotonic()
        writer.execute("BEGIN IMMEDIATE")
        writer.execute("UPDATE companies SET name='Changed' WHERE id='c000001'")
        writer.commit()
        writer.close()
        assert time.monotonic() - start < 0.1
        assert reader.execute("SELECT name FROM companies WHERE id='c000001'").fetchone()[0] == before
        reader.rollback()


def test_no_children_left():
    assert no_child_processes()
