"""Build the legacy v2 fixtures. 只能在 9423c53 代码上运行。

The data-writing part (``--child``) imports ``leadworkbench`` and must resolve to
the 9423c53 code: the driver extracts that commit with ``git archive`` into a
temporary directory and runs itself in a subprocess whose PYTHONPATH points at
it. Outputs (committed under tests/fixtures/):

- legacy_v2_9423c53.sql         full fixture (QA-F1 rejection, accepted-then-rejected,
                                 legacy pending conflict, field review, text_only plan)
- legacy_export_source.sql       ``--export-source``: no repair candidates, no pending conflicts
- legacy_export_9423c53.json     old-code XLSX Contacts/Locations/Sources rows of the export source
- legacy_v2_export_9423c53.json  old-code XLSX rows of the full fixture

SQL files are ``iterdump()`` output; iterdump does not include user_version, so
the loader (tests/legacy_db.py) runs PRAGMA user_version=2 after executescript.
Usage: .venv/bin/python tests/fixtures/build_legacy_v2.py
"""
from __future__ import annotations

from contextlib import closing
import io
import json
import os
from pathlib import Path
import shutil
import sqlite3
import subprocess
import sys
import tempfile

HERE = Path(__file__).resolve().parent
TESTS = HERE.parent


def _listing(n, name, category, address, phone, website):
    from leadworkbench.sources.maps import _record
    raw = dict(title=name, link=f"https://www.google.com/maps/search/?api=1&query_place_id=PLACE_{n}",
               place_id=f"PLACE_{n}", address=address, phone=phone, web_site=website, category=category,
               latitude=37.30 + n / 1000, longitude=-121.88 - n / 1000)
    return _record(raw, f"2026-10-01T0{n % 10}:00:00+00:00")


def _contact(name, role, email, kind, role_status, page, evidence, url=""):
    return dict(name=name, role=role, email=email, url=url, kind=kind, role_status=role_status,
                association_status="same_block" if name else "company_only",
                source_url=page, evidence=evidence)


def _enrich(store, company_id, *contacts, observations=()):
    store.apply_enrichment(company_id, dict(status="completed", contacts=list(contacts),
                                            observations=list(observations)))


def _pending_contact_review(store, contact_id):
    with store.connect() as con:
        return con.execute("SELECT id FROM reviews WHERE kind='contact' AND status='pending' AND new_value=?",
                           (contact_id,)).fetchone()[0]


def _contact_id(store, company_id, email):
    with store.connect() as con:
        return con.execute("SELECT id FROM contacts WHERE company_id=? AND email=?", (company_id, email)).fetchone()[0]


def build_full(directory):
    from leadworkbench.db import Store
    from leadworkbench.market_plan import _job
    store = Store(directory)
    market = store.create_market(dict(country="US", state="California", city="San Jose",
                                      keywords=["metal roofing supply"], company_types=["wholesale"],
                                      auto_enrich=False))
    run = store.start_task(market["id"])
    job = _job("metal roofing supply in San Jose, California, US", None, None, 0)
    store.save_market_plan(run["id"], dict(jobs=[job], coverage={
        "scope_label": "San Jose, California, US · 城市文字检索（范围定位未完成）", "bbox": None,
        "attribution": "", "query_basis": "keywords", "query_terms": ["metal roofing supply"]},
        warnings=["城市范围定位超时（TimeoutError），暂用城市文字检索"], mode="text_only", status="completed"))
    rows = [
        (1, "Alpha Roofing Supply", "Roofing supply store", "100 First St, San Jose, CA 95112", "+1 408-555-0101", "https://alpha.example"),
        (2, "Beta Metals", "Metal supplier", "200 Second St, San Jose, CA 95113", "+1 408-555-0102", "https://beta.example"),
        (3, "Gamma Sheet Metal", "Sheet metal contractor", "300 Third St, Campbell, CA 95008", "+1 408-555-0103", "https://gamma.example"),
        (4, "Delta Building Supply", "Building materials store", "400 Fourth St, San Jose, CA 95110", "+1 408-555-0104", "https://delta.example"),
        (5, "Epsilon Roofing Contractor", "Roofing contractor", "500 Fifth St, San Jose, CA 95111", "", ""),
    ]
    companies = {}
    for n, name, category, address, phone, website in rows:
        companies[n] = store.ingest_market(_listing(n, name, category, address, phone, website),
                                           run["id"], "metal roofing supply", market["id"])["id"]
    a, b, c, d = companies[1], companies[2], companies[3], companies[4]
    # QA-F1 with the old code: Bob Buyer is replaced by Robert Other, then rejected.
    _enrich(store, a, _contact("Bob Buyer", "Buyer", "bob@alpha.example", "procurement", "explicit",
                               "https://alpha.example/team", "Bob Buyer | Buyer | bob@alpha.example"),
            _contact("", "", "info@alpha.example", "general", "unknown", "https://alpha.example/contact",
                     "Email: info@alpha.example"),
            observations=[dict(field="phone", value="+1 408-555-0199", source_url="https://alpha.example/contact",
                               evidence="Phone: +1 408-555-0199")])
    _enrich(store, a, _contact("Robert Other", "Sales Rep", "bob@alpha.example", "general", "explicit",
                               "https://alpha.example/about", "Robert Other | Sales Rep | bob@alpha.example"))
    store.review_decision(_pending_contact_review(store, _contact_id(store, a, "bob@alpha.example")), "reject")
    _enrich(store, b, _contact("Alice Agent", "Purchasing Manager", "alice@beta.example", "procurement", "explicit",
                               "https://beta.example/team", "Alice Agent | Purchasing Manager | alice@beta.example"))
    # Accepted change first, later change rejected: repair must restore 'accepted'.
    _enrich(store, c, _contact("Dana Buyer", "Buyer", "dana@gamma.example", "procurement", "explicit",
                               "https://gamma.example/team", "Dana Buyer | Buyer | dana@gamma.example"))
    _enrich(store, c, _contact("Dana Smith", "Purchasing Director", "dana@gamma.example", "procurement", "explicit",
                               "https://gamma.example/about", "Dana Smith | Purchasing Director | dana@gamma.example"))
    store.review_decision(_pending_contact_review(store, _contact_id(store, c, "dana@gamma.example")), "accept")
    _enrich(store, c, _contact("Dee Other", "Sales Rep", "dana@gamma.example", "general", "explicit",
                               "https://gamma.example/news", "Dee Other | Sales Rep | dana@gamma.example"))
    store.review_decision(_pending_contact_review(store, _contact_id(store, c, "dana@gamma.example")), "reject")
    # Legacy pending conflict, left undecided.
    _enrich(store, d, _contact("Carl Buyer", "Buyer", "carl@delta.example", "procurement", "explicit",
                               "https://delta.example/team", "Carl Buyer | Buyer | carl@delta.example"))
    _enrich(store, d, _contact("Carla Other", "Sales Rep", "carl@delta.example", "general", "explicit",
                               "https://delta.example/about", "Carla Other | Sales Rep | carl@delta.example"))
    store.finalize(run["id"])
    return store


def build_export_source(directory):
    from leadworkbench.db import Store
    store = Store(directory)
    market = store.create_market(dict(country="US", state="California", city="San Jose",
                                      keywords=["roofing"], company_types=["wholesale"], auto_enrich=False))
    run = store.start_task(market["id"])
    rows = [
        (11, "Echo Roofing Wholesale", "Roofing supply store", "11 Echo Ave, San Jose, CA 95112", "+1 408-555-0111", "https://echo.example"),
        (12, "Foxtrot Metals", "Metal supplier", "12 Fox Rd, Campbell, CA 95008", "+1 408-555-0112", "https://fox.example"),
        (13, "Golf Refusal Co", "Roofing supply store", "13 Golf St, San Jose, CA 95110", "+1 408-555-0113", "https://golf.example"),
    ]
    ids = {n: store.ingest_market(_listing(n, *rest), run["id"], "roofing", market["id"])["id"] for n, *rest in rows}
    e, f, g = ids[11], ids[12], ids[13]
    # Same mailbox on two pages: one entry, two sources.
    _enrich(store, e, _contact("", "", "info@echo.example", "general", "unknown", "https://echo.example/contact",
                               "Email: info@echo.example"))
    _enrich(store, e, _contact("", "", "info@echo.example", "general", "unknown", "https://echo.example/about",
                               "Write to info@echo.example"),
            _contact("Erin Buyer", "Buyer", "erin@echo.example", "procurement", "explicit", "https://echo.example/team",
                     "Erin Buyer | Buyer | Phone: +1 408 555 0177 | erin@echo.example"),
            _contact("", "", "", "contact_form", "unknown", "https://echo.example/contact",
                     "Contact us form", url="https://echo.example/contact"))
    # Alex then Bob share one mailbox; the accepted change keeps Alex's version separate.
    _enrich(store, f, _contact("Alex Seller", "Sales Manager", "team@fox.example", "general", "explicit",
                               "https://fox.example/team", "Alex Seller | Sales Manager | Phone: +1 408 555 0188 | team@fox.example"))
    _enrich(store, f, _contact("Bob Buyer", "Buyer", "team@fox.example", "procurement", "explicit",
                               "https://fox.example/about", "Bob Buyer | Buyer | team@fox.example"))
    with store.connect() as con:
        review = con.execute("SELECT id FROM reviews WHERE kind='contact' AND status='pending'").fetchone()[0]
    store.review_decision(review, "accept")
    _enrich(store, g, _contact("Gail Buyer", "Buyer", "gail@golf.example", "procurement", "explicit",
                               "https://golf.example/team", "Gail Buyer | Buyer | gail@golf.example"))
    company = store.company(g)
    store.update_company(g, company["revision"], {}, status="do_not_contact")
    store.finalize(run["id"])
    with store.connect() as con:
        assert not con.execute("SELECT 1 FROM reviews WHERE status='pending' AND kind='contact'").fetchone()
    return store


def export_rows(store):
    from openpyxl import load_workbook
    from leadworkbench.export import export_data
    book = load_workbook(io.BytesIO(export_data(store, "xlsx")))
    return {name: [[cell if cell is not None else "" for cell in row] for row in book[name].iter_rows(values_only=True)]
            for name in ("Contacts", "Locations", "Sources")}


def dump(store, target):
    with closing(sqlite3.connect(store.path)) as con:
        lines = list(con.iterdump())
    header = ("-- Generated by tests/fixtures/build_legacy_v2.py with leadworkbench @ 9423c53. Synthetic data only.\n"
              "-- iterdump 不包含 user_version；加载器在 executescript 之后执行 PRAGMA user_version=2\n")
    target.write_text(header + "\n".join(lines) + "\n", encoding="utf-8")


def child(out):
    import leadworkbench
    assert "legacy-code-" in leadworkbench.__file__, leadworkbench.__file__
    out = Path(out)
    with tempfile.TemporaryDirectory() as tmp:
        full = build_full(Path(tmp) / "full")
        dump(full, out / "legacy_v2_9423c53.sql")
        (out / "legacy_v2_export_9423c53.json").write_text(
            json.dumps(export_rows(full), ensure_ascii=False, indent=1), encoding="utf-8")
        source = build_export_source(Path(tmp) / "source")
        dump(source, out / "legacy_export_source.sql")
        (out / "legacy_export_9423c53.json").write_text(
            json.dumps(export_rows(source), ensure_ascii=False, indent=1), encoding="utf-8")


def main():
    if "--child" in sys.argv:
        return child(sys.argv[sys.argv.index("--child") + 1])
    sys.path.insert(0, str(TESTS))
    from legacy_db import legacy_code_tmp
    code = legacy_code_tmp()
    if code is None:
        raise SystemExit("git is required to extract 9423c53")
    try:
        env = {**os.environ, "PYTHONPATH": str(code), "LEAD_DATA_DIR": tempfile.mkdtemp()}
        env.pop("MAPS_SCRAPER_BIN", None)
        subprocess.run([sys.executable, str(Path(__file__).resolve()), "--child", str(HERE)],
                       env=env, cwd=code, check=True)
    finally:
        shutil.rmtree(code, ignore_errors=True)
        shutil.rmtree(env["LEAD_DATA_DIR"], ignore_errors=True)


if __name__ == "__main__":
    main()
