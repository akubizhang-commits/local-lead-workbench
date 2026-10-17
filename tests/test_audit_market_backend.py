"""F1/F5/F8/F9/U5 regressions: isolated stores, no network or production data."""
import csv
import io

import pytest
from openpyxl import load_workbook

from leadworkbench.db import Store
from leadworkbench.export import export_data
from leadworkbench.market_store import MarketStoreMixin
from leadworkbench.sources.maps import _record


@pytest.fixture
def setup(tmp_path):
    store = Store(tmp_path)
    market = store.create_market(dict(name='Audit',country='US',state='California',city='Los Angeles',keywords=['retail'],company_types=[],auto_enrich=False))
    return store, market, store.start_task(market['id'])


def listing(place='PLACE_A', cid=None, address='1 Old Street Los Angeles', phone='12125550111'):
    raw = dict(title='Fixture Brand',link=f'https://www.google.com/maps/search/?api=1&query_place_id={place}',place_id=place,address=address,phone=phone,web_site='https://fixture.example')
    if cid:
        raw['cid'] = cid
    return _record(raw, '2026-09-22T01:00:00+00:00')


def ingest(setup, record):
    s,m,r=setup
    return s.ingest_market(record,r['id'],'fixture',m['id'])


def test_f1_place_alias_move_keeps_one_location_and_contacts(setup):
    s,m,r=setup
    a=ingest(setup,listing())
    s.apply_enrichment(a['id'],dict(status='completed',contacts=[dict(name='Pat',email='pat@fixture.example',source_url='https://fixture.example/team',evidence='Public card')]))
    b=ingest(setup,listing(cid='123',address='2 New Street Los Angeles',phone='12125550222'))
    assert b['id']==a['id']
    assert s.detail(b['id'])['contacts'][0]['email']=='pat@fixture.example'
    assert len(s.companies())==1
    with s.connect() as con:
        assert con.execute('SELECT count(*) FROM locations').fetchone()[0]==1
        assert con.execute('SELECT address FROM locations').fetchone()[0]=='2 New Street Los Angeles'
    assert any(x['field']=='address' and x['old_value']=='1 Old Street Los Angeles' for x in s.changes(a['id']))
    branch=ingest(setup,listing('PLACE_B','456','3 Branch Street Los Angeles','12125550333'))
    assert branch['id'] != a['id']


def test_f1_backfills_legacy_listing_aliases_before_move(setup):
    s,m,r=setup
    a=listing()
    # Legacy ingestion has no alias table knowledge.
    c=s.ingest(a,r['id'],'legacy',city='Los Angeles',country='US')
    b=ingest(setup,listing(cid='123',address='2 New Street Los Angeles'))
    assert b['id']==c['id']
    with s.connect() as con:
        s._migrate_source_aliases(con)
        s._migrate_source_aliases(con)
        assert con.execute('SELECT count(*) FROM locations').fetchone()[0]==1


def test_f1_conflicting_known_cids_do_not_merge_branches(setup):
    s,m,r=setup
    a=ingest(setup,listing(cid='123'))
    b=ingest(setup,listing('PLACE_B','456','3 Branch Street Los Angeles','12125550333'))
    with pytest.raises(ValueError, match='标识|CID'):
        ingest(setup,listing('PLACE_A','456','3 Branch Street Los Angeles','12125550333'))
    assert len(s.companies())==2
    assert s.company(a['id'])['address']=='1 Old Street Los Angeles'
    assert s.company(b['id'])['address']=='3 Branch Street Los Angeles'


def test_f1_delayed_observation_can_teach_alias_without_reverting_move(setup):
    s,m,r=setup
    c=ingest(setup,listing('PLACE_A','123','2 New Street Los Angeles'))
    delayed={**listing('PLACE_NEW','123'), 'observed_at':'2026-09-21T01:00:00+00:00'}
    ingest(setup,delayed)
    later={**listing('PLACE_NEW',None,'3 Latest Street Los Angeles'), 'observed_at':'2026-09-23T01:00:00+00:00'}
    assert ingest(setup,later)['id']==c['id']
    assert len(s.companies())==1


def test_f1_changed_accounting_uses_alias_identity_when_listing_url_changes(setup):
    s,m,r=setup
    c=ingest(setup,listing())
    with s.connect() as con:
        con.execute("UPDATE work_items SET state='succeeded' WHERE run_id=?", (r['id'],))
    s.finalize(r['id'])
    next_run=s.start_task(m['id'])
    raw=dict(title='Fixture Brand',link='https://www.google.com/maps?cid=123',
             place_id='PLACE_A',cid='123',address='2 New Street Los Angeles',phone='12125550111',
             web_site='https://fixture.example')
    moved=s.ingest_market(_record(raw,'2026-09-23T01:00:00+00:00'),next_run['id'],'fixture',m['id'])
    assert moved['id']==c['id'] and moved['_outcome']=='existing'
    assert moved['_changed'] is True


@pytest.mark.asyncio
@pytest.mark.parametrize('bbox,level',[(None,0),([34,35,-119,-118],3)])
async def test_f5_exhaustion_is_durable_and_cannot_call_source_on_replay(setup,bbox,level):
    from leadworkbench.worker import process_run
    s,m,r=setup
    async def plan(*args):
        return dict(status='completed',warnings=[],coverage={},jobs=[dict(key='leaf',query='fixture',bbox=bbox,geo=None,zoom=15,depth=10,level=level)])
    async def source(*args,**kwargs):
        return dict(status='completed',records=[listing()],saturated=True,warnings=[],stop_reason='company_limit')
    await process_run(s,r['id'],planner=plan,discover=source)
    coverage=s.task(m['id'])['coverage']
    assert coverage.get('retryable') is False
    assert coverage.get('capability_exhausted') is True
    with s.connect() as con:
        assert s.market_resume_allowed(con,r['id']) is False
        item=dict(con.execute("SELECT * FROM work_items WHERE stage='discovery'").fetchone())
    # A stale client/worker replay must still avoid another ineffective source call.
    async def forbidden(*args,**kwargs):
        pytest.fail('capability exhausted work cannot replay the source')
    from leadworkbench.market_worker import discover_market_item
    result=await discover_market_item(s,s.run(r['id']),item,forbidden)
    assert result['status']=='partial' and result['retryable'] is False


def test_f5_resume_contract_preserves_exhausted_leaves_but_retries_other_work(setup):
    s,m,r=setup
    s.save_market_plan(r['id'],dict(status='completed',jobs=[dict(key='a',query='a'),dict(key='b',query='b')]))
    with s.connect() as con:
        con.execute("UPDATE work_items SET state='succeeded' WHERE stage='planning'")
        items=[dict(row) for row in con.execute("SELECT * FROM work_items WHERE stage='discovery' ORDER BY rowid")]
    for item in items:
        s.finish_item(item['id'],'partial',{})
    s.record_market_exhausted(r['id'],items[0]['id'],'capability exhausted')
    with s.connect() as con:
        assert s.market_resume_allowed(con,r['id']) is True
        assert s.market_retryable_items(con,r['id'])==[items[1]['id']]
    s.finish_item(items[1]['id'],'succeeded',{})
    s.finalize(r['id'])
    assert s.run(r['id'])['status']=='partial'
    assert s.task(m['id'])['coverage']['retryable'] is False


@pytest.mark.parametrize('term',['California','US','Old Street'])
def test_f8_region_search_list_and_export(setup,term):
    s,m,r=setup
    c=ingest(setup,listing())
    assert [v['id'] for v in s.company_page(q=term)['items']]==[c['id']]
    rows=list(csv.DictReader(io.StringIO(export_data(s,q=term).decode('utf-8-sig'))))
    assert [v['id'] for v in rows]==[c['id']]


@pytest.mark.parametrize('format',['csv','xlsx'])
@pytest.mark.parametrize('empty',[False,True])
def test_f9_run_only_export_has_market_metadata_even_when_empty(setup,format,empty):
    s,m,r=setup
    ingest(setup,listing())
    def rows(**filters):
        raw=export_data(s,format,**filters)
        return list(csv.reader(io.StringIO(raw.decode('utf-8-sig')))) if format=='csv' else list(load_workbook(io.BytesIO(raw))['Companies'].values)
    filters={'q':'missing'} if empty else {}
    direct=rows(run_id=r['id'],**filters)
    explicit=rows(run_id=r['id'],market_id=m['id'],**filters)
    assert direct==explicit
    assert set(['market_segment','market_changed','first_seen','last_seen']) <= set(direct[0])


def add_contact(s,c,name,kind='general',role_status='explicit',email=True,exclude=False,reject=False):
    s.apply_enrichment(c['id'],dict(status='completed',contacts=[dict(name=name,email=f'{name.lower()}@fixture.example' if email else '',role='Buyer' if kind=='procurement' else 'Director',url='https://fixture.example/contact',kind=kind,role_status=role_status,source_url='https://fixture.example/team',evidence='Public card')]))
    with s.connect() as con:
        cid=con.execute('SELECT id FROM contacts WHERE company_id=? AND name=?',(c['id'],name)).fetchone()[0]
        if exclude:
            con.execute("INSERT INTO contact_decisions VALUES(?,'excluded','2026-09-22')",(cid,))
        if reject:
            con.execute("UPDATE contacts SET review_status='reject' WHERE id=?",(cid,))
    return cid


def test_u5_usable_contact_filter_counts_sort_and_export(setup):
    s,m,r=setup
    buyer=s.manual_company('Buyer','https://buyer.example')
    mail=s.manual_company('Mail','https://mail.example')
    excluded=s.manual_company('Excluded','https://excluded.example')
    add_contact(s,buyer,'Pat','procurement')
    add_contact(s,buyer,'Alex','procurement','historical',email=False)
    add_contact(s,mail,'Chris');add_contact(s,mail,'Sam')
    add_contact(s,excluded,'Robin','procurement',reject=True)
    add_contact(s,excluded,'Taylor','procurement',exclude=True)
    page=s.company_page(contact='email',sort='procurement',page_size=1)
    assert page['total']==1 and page['items'][0]['id']==mail['id']
    assert page['items'][0]['email_count']==2 and page['items'][0]['procurement_count']==0
    assert [c['id'] for c in s.company_page(contact='procurement')['items']]==[buyer['id']]
    assert s.company_page(sort='email')['items'][0]['id']==mail['id']
    # Mixin is the owned surface; parent forwards these keywords on Store.companies.
    assert MarketStoreMixin.companies(s,contact='email',sort='email')[0]['id']==mail['id']
    rows=list(csv.DictReader(io.StringIO(export_data(s,contact='email',sort='email').decode('utf-8-sig'))))
    assert [row['id'] for row in rows]==[mail['id']]
    contacts=list(load_workbook(io.BytesIO(export_data(s,'xlsx')))['Contacts'].values)
    assert not any(row[1] in {'Taylor','Robin'} for row in contacts[1:])


@pytest.mark.parametrize('state',['pending','reject','invalid','superseded'])
def test_u5_unusable_contact_states_never_count_filter_or_export(setup,state):
    s,m,r=setup
    c=s.manual_company('Fixture','https://fixture.example')
    cid=add_contact(s,c,'Pat','procurement')
    with s.connect() as con:
        con.execute('UPDATE contacts SET review_status=? WHERE id=?',(state,cid))
    page=s.company_page()
    assert page['items'][0]['email_count']==0
    assert page['items'][0]['procurement_count']==0
    assert s.company_page(contact='email')['total']==0
    assert s.company_page(contact='procurement')['total']==0
    assert len(list(load_workbook(io.BytesIO(export_data(s,'xlsx')))['Contacts'].values))==1


def test_u5_contact_counts_deduplicate_same_person_across_sources(setup):
    s,m,r=setup
    c=s.manual_company('Fixture','https://fixture.example')
    cid=add_contact(s,c,'Pat','procurement')
    # Existing databases can contain one row per source; preserve both evidence rows.
    with s.connect() as con:
        row=dict(con.execute('SELECT * FROM contacts WHERE id=?',(cid,)).fetchone())
        row.update(id='second-source',source_url='https://fixture.example/about',fingerprint='second-source')
        con.execute('INSERT INTO contacts VALUES('+','.join('?' for _ in row)+')',tuple(row.values()))
        assert con.execute('SELECT count(*) FROM contacts WHERE company_id=?',(c['id'],)).fetchone()[0]==2
    page=s.company_page(contact='procurement')
    assert page['items'][0]['email_count']==0
    assert page['items'][0]['procurement_count']==1
    assert page['items'][0]['contact_count']==1
    contacts=list(load_workbook(io.BytesIO(export_data(s,'xlsx')))['Contacts'].values)
    assert len(contacts)==2


def test_u5_updated_sort_and_invalid_contract(setup):
    s,m,r=setup
    old=s.manual_company('Old','https://old.example')
    new=s.manual_company('New','https://new.example')
    with s.connect() as con:
        con.execute("UPDATE companies SET created_at='2020-01-01',last_checked_at='2026-09-23' WHERE id=?",(old['id'],))
        con.execute("UPDATE companies SET created_at='2026-01-01',last_checked_at='2026-09-22' WHERE id=?",(new['id'],))
    assert s.company_page(sort='newest')['items'][0]['id']==new['id']
    assert s.company_page(sort='updated')['items'][0]['id']==old['id']
    with pytest.raises(ValueError):
        s.company_page(contact='unknown')
    with pytest.raises(ValueError):
        s.company_page(sort='unknown')
