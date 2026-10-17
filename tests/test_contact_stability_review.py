"""Regression cases from the September contact audit; no network or live DB."""
import io

from openpyxl import load_workbook

from leadworkbench.contact_projection import group_contact_records, outreach_category
from leadworkbench.db import Store
from leadworkbench.export import export_data


def record(email):
    return dict(name='', role='', email=email, url='', kind='general',
                role_status='unknown', association_status='company_only',
                source_url='https://brand.example/contact', evidence=f'Email: {email}')


def test_placeholder_and_malformed_mail_are_not_resalvaged(tmp_path):
    store = Store(tmp_path)
    company = store.manual_company('Audit fixture', 'https://brand.example')
    emails = ['youremail@here.com', 'filler@godaddy.com', 'a..b@brand.example.com',
              'hello@@brand.example.com', 'mail@-brand.example.com', 'mail@brand..example.com']
    store.apply_enrichment(company['id'], {'status': 'completed',
                                         'contacts': [record(e) for e in emails]})
    assert store.detail(company['id'])['contacts'] == []
    assert store.company(company['id'])['contact_count'] == 0
    assert store.reclassify_contacts()['archived'] == 0
    with store.connect() as con:
        assert con.execute('SELECT count(*) FROM contacts').fetchone()[0] == len(emails)
        assert all(r[0] == 'invalid' for r in con.execute('SELECT review_status FROM contacts'))


def test_non_outreach_contacts_are_not_counted_or_exported(tmp_path):
    store = Store(tmp_path)
    company = store.manual_company('Audit fixture', 'https://brand.example')
    emails = ['abuse@company.site.example.org', 'Privacy@vestis.example.org', 'data.privacy@dssmith.example.org',
              'My-Privacy@ipaper.example.org', 'no-reply@brand.example.com', 'jobs@brand.example.com',
              'sales@brand.example.com', 'support@brand.example.com', 'hello@brand.example.com']
    store.apply_enrichment(company['id'], {'status': 'completed',
                                         'contacts': [record(e) for e in emails]})
    company_row = store.company_page()['items'][0]
    assert company_row['contact_count'] == company_row['email_count'] == 3
    assert store.company(company['id'])['contact_count'] == 3
    contacts = store.detail(company['id'])['contacts']
    assert sum(r['outreach_category'] == 'irrelevant' for r in contacts) == 6
    book = load_workbook(io.BytesIO(export_data(store, 'xlsx')))
    assert book['Contacts'].max_row == 4
    # Keep the captured information and provenance available for inspection.
    with store.connect() as con:
        assert con.execute('SELECT count(*) FROM contacts').fetchone()[0] == 9


def test_old_or_accepted_placeholder_cannot_leak_into_display():
    for status in ['auto_recorded', 'pending', 'accept']:
        row = dict(id='legacy', review_status=status, **record('youremail@here.com'))
        assert group_contact_records([row], include_excluded=True) == []
        assert outreach_category('procurement', 'explicit', row['email'], 'Buyer') == 'irrelevant'


def test_channel_rules_do_not_hide_plausible_commercial_mailboxes():
    for email in ['sales@brand.example.com', 'support@brand.example.com', 'care@brand.example.com',
                  'privacyrose@gmail.example.org', 'chris@brand.example.com', 'hello@jobs-fashion.example.org',
                  'alice+buying@brand.example.com', "o'neill@brand.example.com", 'info@品牌.example.org']:
        assert outreach_category('general', 'unknown', email, '') == 'email'
    assert outreach_category('procurement', 'explicit', 'buyer@brand.example.com', 'Buyer') == 'procurement'


def test_placeholder_navigation_and_subscription_forms_are_not_channels(tmp_path):
    store = Store(tmp_path)
    company = store.manual_company('Audit fixture', 'https://brand.example')
    false_forms = [
        ('https://brand.example/about', ' | Contact Us | #'),
        ('https://brand.example/newsletter-sign-up/', 'First name Email Contact number Location Accept terms Yes I agree'),
        ('https://brand.example/ProductResults/?ProdSetIds=158', 'Office Location Phone: 2132221111 Search all Products Home About Us Contact Us'),
        ('https://brand.example/', 'Office Location Phone: 2132221111 Search all Products Price From $ - Price To $ Go Home About Us Contact Us'),
        ('https://brand.example/default.aspx', ' | Home Bulls Cows Merch About Contact Welcome to our ranch Privacy Policy Login'
         ' | form action=https://brand.example/default.aspx | / | /Bulls | /Cows | /Merch | /About | /Contact'),
    ]
    forms = [{**record(''), 'url': url, 'evidence': evidence, 'kind': 'contact_form'}
             for url, evidence in false_forms]
    forms.append({**record(''), 'kind': 'contact_form', 'url': 'https://brand.example/contact',
                  'evidence': 'Contact us | /contact'})
    store.apply_enrichment(company['id'], {'status': 'completed', 'contacts': forms})
    assert len(store.detail(company['id'])['contacts']) == 1
    assert store.company_page()['items'][0]['contact_count'] == 1
    # Simulate legacy rows not yet reclassified; the read path must agree.
    with store.connect() as con:
        con.execute("UPDATE contacts SET review_status='accept'")
    assert len(store.detail(company['id'])['contacts']) == 1
    assert store.company(company['id'])['contact_count'] == 1
    assert store.company_page()['items'][0]['contact_count'] == 1


def test_contact_links_merge_tracking_variants_but_preserve_distinct_destinations():
    base = {**record(''), 'kind': 'contact_form', 'review_status': 'auto_recorded',
            'url': 'https://brand.example/ContactUs/?referrerPage=Home&refPgId=1',
            'evidence': 'Contact Us'}
    result = group_contact_records([
        {**base, 'id': 'a'},
        {**base, 'id': 'b', 'url': 'https://brand.example/ContactUs/?referrerPage=AboutUs&refPgId=2'},
        {**base, 'id': 'c', 'url': 'https://brand.example/ContactUs/?department=buying'},
    ])
    assert len(result) == 2
    assert len(result[0]['record_ids']) == 2


def test_surname_jobs_is_not_a_recruitment_mailbox(tmp_path):
    store = Store(tmp_path)
    company = store.manual_company('Audit fixture', 'https://brand.example')
    person = {**record('steve.jobs@brand.example'), 'name': 'Steve Jobs', 'role': 'Buyer',
              'kind': 'procurement', 'role_status': 'explicit', 'association_status': 'same_block'}
    store.apply_enrichment(company['id'], {'status': 'completed', 'contacts': [person]})
    with store.connect() as con:
        con.execute("UPDATE contacts SET review_status='accept'")
    assert store.company_page()['items'][0]['procurement_count'] == 1
    book = load_workbook(io.BytesIO(export_data(store, 'xlsx')))
    assert book['Contacts'].max_row == 2


def test_real_message_form_survives_product_search_copy():
    from leadworkbench.sources.website import parse_page
    html = """<form action="/help"><h2>Contact us</h2>
      <p>Search all products, then send us your question.</p>
      <label for="message">Message</label><textarea id="message" name="message"></textarea>
      <button>Send message</button></form>"""
    contacts = parse_page(html, 'https://brand.example/help')['contacts']
    assert len(contacts) == 1 and contacts[0]['kind'] == 'contact_form'


def test_missing_mailbox_does_not_promote_an_unrelated_form_to_procurement(tmp_path):
    store = Store(tmp_path)
    company = store.manual_company('Audit fixture', 'https://brand.example')
    buyer = {**record(''), 'name': 'Alice Brown', 'role': 'Buyer',
             'kind': 'procurement', 'role_status': 'explicit'}
    form = {**record(''), 'url': 'https://brand.example/contact',
            'kind': 'contact_form', 'evidence': 'Contact us'}
    store.apply_enrichment(company['id'], {'status': 'completed', 'contacts': [buyer, form]})
    contacts = store.detail(company['id'])['contacts']
    assert sum(row['outreach_category'] == 'procurement' for row in contacts) == 1
    assert store.company_page()['items'][0]['procurement_count'] == 1
